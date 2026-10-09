"""Builds Foxy_Fight_Sheet_Template.xlsx from data/*.csv.
Upload it to Google Drive and open with Google Sheets (or File > Import).
Every data cell is stored as plain text so Google's CSV export never drops values
in columns that mix numbers and words."""
import csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.join(os.path.dirname(__file__), '..')
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'dist', 'Foxy_Fight_Sheet_Template.xlsx')
TABS = ['Config', 'Styles', 'Cards', 'Statuses', 'Opponents', 'CutIns', 'Fighters']
OPTIONAL = {'CutIns', 'Fighters'}
HEAD_FILL = {'Config': 'F3E2B8', 'Styles': 'F2C9D1', 'Cards': 'EBD8AE', 'Statuses': 'CFE6E0', 'Opponents': 'E6D3EE',
             'CutIns': 'F6D6B8', 'Fighters': 'D8E2F2'}
WIDTHS = {'description': 60, 'effects': 26, 'link_bonus': 26, 'base_deck': 70, 'deck': 70, 'text': 34, 'flavor': 34,
          'intro': 40, 'win_art': 22, 'loss_art': 22, 'rounds': 9, 'notes': 70, 'follows': 20, 'name': 20, 'title': 18, 'on_turn_start': 18, 'on_link': 18,
          'fighter': 18, 'move': 14, 'profile_0': 46, 'profile_1': 46, 'profile_2': 46, 'pov_0': 46, 'pov_1': 46, 'pov_2': 46,
          'outfit_of': 14, 'style': 16}

wb = Workbook()
guide = wb.active
guide.title = 'How to use'
lines = [
    ('Foxy Fight game data', True),
    ('Each tab below is read by the game. Keep the tab names and the header row exactly as they are; add columns of your own to the right if you like (the game ignores them).', False),
    ('Rows whose first cell starts with # are ignored, so you can use them for notes or to switch a card off.', False),
    ('The source column marks who wrote a row: set it to mine on rows you write or edit, and Claude will leave that text alone.', False),
    ('', False),
    ('Connect it to the game', True),
    ('1. Share > General access > Anyone with the link > Viewer.', False),
    ('2. In the game, open Sheet & saves, paste this sheet\'s link and press Sync from sheet.', False),
    ('3. If something is wrong, the game lists the problem with its tab and row and keeps the previous data.', False),
    ('', False),
    ('Effects (cards: effects, link_bonus; styles: on_turn_start, on_link). Separate several with ;', True),
    ('dmg:5   deal 5 damage        dmg:2x3   three hits of 2', False),
    ('pierce:4   damage that ignores block', False),
    ('dmgchain:2   2 damage for each card in the current chain, as one hit', False),
    ('block:5   heal:4   draw:1   energy:1', False),
    ('apply:bleed:2   give the opponent a status        gain:focus:2   give yourself a status', False),
    ('Effects use the status id, not its display name: bleed = Cut, exposed = Wide Open, focus = Swagger, stagger = Winded, dazed = Seeing Stars, fortify = Tight Guard, thorns = Counterpunch.', False),
    ('', False),
    ('Chains', True),
    ('tags: what the card is (Opener, Light, Heavy, Finisher, Guard, Wild, or any new tag you invent).', False),
    ('follows: the tags this card can come after. START = first card of the turn. ANY = after any card.', False),
    ('A card tagged Wild lets any card link after it.', False),
    ('', False),
    ('Rounds', True),
    ('Config: exchanges_per_round (0 = no rounds), rounds, corner_heal_percent, corner_clears_statuses (stacks removed, or TRUE/FALSE), corner_reshuffle, alternate_round_start, knockdowns_for_tko, knockdown_recover_percent, knockdown_clears_statuses.', False),
    ('Opponents: rounds overrides the Config value for that fight (blank = Config).', False),
    ('', False),
    ('Versus', True),
    ('Config: pvp_turn_seconds (turn clock, 0 = none), pvp_disconnect_seconds (wait before a forfeit), pvp_xp_win, pvp_xp_loss.', False),
    ('', False),
    ('Decks and upgrades', True),
    ('Deck lists look like: ps_jab x3, ps_kick x2, ps_flurry', False),
    ('upgrade_of: buying this card turns one owned copy of that card into this one.', False),
    ('style: which style can buy the card. Use Neutral for cards every style can buy. Cards of an evolved style unlock when you evolve.', False),
    ('xp_price: leave blank to keep a card out of the shop.', False),
    ('', False),
    ('Statuses (behaviors you can reuse for new statuses)', True),
    ('dot, damage_taken_up, damage_dealt_up, energy_down, draw_down, block_gain_up, thorns', False),
    ('decay: turn (lose 1 stack at the start of the owner\'s turn), clear (remove at the start of the owner\'s turn), consume (used up by the next attack), none', False),
    ('', False),
    ('Art', True),
    ('Put an image filename (e.g. ps_jab.png, served from the game\'s art/ folder) or a full https:// image link in the art column. Blank = generated art-deco art. A value with a / in it (art/cutins/...) is used as a path from the game folder.', False),
    ('End-of-fight images: Opponents tab win_art (shown when you beat her) and loss_art (shown when she beats you). Config default_win_art / default_loss_art cover any blanks. Wide 16:9 images work best.', False),
    ('', False),
    ('Cut-in art (CutIns and Fighters tabs)', True),
    ('You don\'t need to edit anything here to add art. Every picture already has its path in CutIns. Drop the files into the game\'s art/cutins/<fighter>/ folders with the names listed in each folder\'s README.txt and upload. A picture that isn\'t there yet is skipped and the nearest one is used instead.', False),
    ('If your own Google Sheet has no CutIns or Fighters tab, the game uses its built-in copy. Add these tabs to your sheet only if you want to change paths.', False),
    ('CutIns: one row per fighter + move. fighter = a style id, an opponent id or a Fighters id. move = a card id, a tag (Opener, Light, Heavy, Finisher, Guard, Wild), Knockdown, Victory, or blank for any move. profile_0..2 / pov_0..2 = image paths or https:// links.', False),
    ('The numbers: on card and tag rows, knockdowns that fighter has taken (0 fresh, 1 worn, 2 battered). Knockdown row: 0 = first knockdown, 1 = second, 2 = T.K.O. Victory row: the knockdowns the winner took.', False),
    ('Your fighter is always shown in profile, the opponent in POV. Ladder opponents only have pov paths. Upgrades (+) use their base card\'s row; Finishers and Elites have their own rows.', False),
    ('Fighters: custom looks a player can pick. style = the style she dresses. outfit_of = the Fighters id this is an evolved outfit of (blank for her base look). Her rows use tags, not cards. Run npm run art-list after adding one to get her paths and folder.', False),
    ('Art problems never block the sheet: a typo shows as a warning and that picture falls back.', False),
]
for i, (t, bold) in enumerate(lines, start=1):
    c = guide.cell(row=i, column=1, value=t)
    c.font = Font(bold=bold, size=13 if bold and i == 1 else 11)
    c.alignment = Alignment(wrap_text=True, vertical='top')
guide.column_dimensions['A'].width = 120

def text_tab(ws, rows, fill):
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.number_format = '@'
            cell.alignment = Alignment(vertical='top', wrap_text=(len(val) > 40))
            if r == 1:
                cell.font = Font(bold=True)
                cell.fill = PatternFill('solid', fgColor=fill)

for tab in TABS:
    path = os.path.join(ROOT, 'data', tab + '.csv')
    if tab in OPTIONAL and not os.path.exists(path):
        continue
    ws = wb.create_sheet(tab)
    with open(path, newline='', encoding='utf-8') as f:
        rows = list(csv.reader(f))
    text_tab(ws, rows, HEAD_FILL[tab])
    # pre-format extra empty rows as text too, so new entries stay text
    for r in range(len(rows) + 1, len(rows) + 60):
        for c in range(1, len(rows[0]) + 1):
            ws.cell(row=r, column=c).number_format = '@'
    for c, h in enumerate(rows[0], start=1):
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS.get(h, max(10, len(h) + 4))
    ws.freeze_panes = 'C2' if tab == 'CutIns' else 'B2'

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('Saved', OUT)
